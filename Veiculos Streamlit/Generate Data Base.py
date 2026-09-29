import pandas as pd
import random
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# raw data lists
first_names_m = ['James', 'John', 'Robert', 'Michael', 'William', 'David', 'Richard', 'Joseph', 'Thomas', 'Charles', 'Daniel', 'Matthew']
first_names_f = ['Mary', 'Patricia', 'Jennifer', 'Linda', 'Elizabeth', 'Barbara', 'Susan', 'Jessica', 'Sarah', 'Karen', 'Nancy', 'Lisa']
last_names = ['Smith', 'Johnson', 'Williams', 'Brown', 'Jones', 'Garcia', 'Miller', 'Davis', 'Rodriguez', 'Martinez', 'Hernandez', 'Lopez']

# map cities to states so we dont mix them up
locations = {
    'CA': ['Los Angeles', 'San Francisco', 'San Diego', 'Sacramento'],
    'NY': ['New York', 'Buffalo', 'Rochester', 'Yonkers'],
    'TX': ['Houston', 'San Antonio', 'Dallas', 'Austin'],
    'FL': ['Miami', 'Orlando', 'Tampa', 'Jacksonville'],
    'IL': ['Chicago', 'Aurora', 'Naperville', 'Joliet']
}

vehicles = ['Toyota Camry', 'Honda Civic', 'Ford F-150', 'Chevrolet Silverado', 'Tesla Model 3', 'Nissan Altima', 'Jeep Wrangler']
car_parts = ['Engine', 'Transmission', 'Brake Pads', 'Shock Absorber', 'Oil Filter', 'Battery', 'Tire', 'Radiator']
education = ['High School', 'Some College', 'Bachelors', 'Masters', 'PhD']
marital_status = ['Single', 'Married', 'Divorced', 'Widowed']
neighborhoods = ['Downtown', 'Westside', 'Eastside', 'North End', 'South Side', 'Midtown', 'Uptown']
streets = ['Main St', 'Oak St', 'Maple Ave', 'Cedar Ln', 'Elm St', 'Washington Blvd', 'Lake St']

data = []

# gen 400 rows
for i in range(1, 401):
    gender = random.choice(['M', 'F'])
    
    # match name to gender
    first_name = random.choice(first_names_m) if gender == 'M' else random.choice(first_names_f)
    last_name = random.choice(last_names)
    
    age = random.randint(18, 80)
    state = random.choice(list(locations.keys()))
    city = random.choice(locations[state])
    neighborhood = random.choice(neighborhoods)
    street = random.choice(streets)
    building_num = str(random.randint(10, 9999))
    
    vehicle = random.choice(vehicles)
    part = random.choice(car_parts)
    part_price = round(random.uniform(50.0, 4500.0), 2)
    income = round(random.uniform(2000.0, 20000.0), 2)
    
    degree = random.choice(education)
    
    # weight kids distribution
    kids = random.choices([0, 1, 2, 3, 4, 5], weights=[40, 25, 20, 10, 4, 1])[0]
    status = random.choice(marital_status)
    
    data.append({
        'ID': i,
        'First Name': first_name,
        'Last Name': last_name,
        'Age': age,
        'Gender': gender,
        'Street': street,
        'Building Number': building_num,
        'Neighborhood': neighborhood,
        'City': city,
        'State': state,
        'Education': degree,
        'Marital Status': status,
        'Kids': kids,
        'Monthly Income (USD)': income,
        'Vehicle': vehicle,
        'Car Part': part,
        'Part Price (USD)': part_price
    })

# create df
df = pd.DataFrame(data)
file_path = 'mock_database.xlsx'

# dump to excel
df.to_excel(file_path, index=False, engine='openpyxl')

# basic excel styling
wb = load_workbook(file_path)
ws = wb.active
ws.title = 'Dataset'

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)

# auto adjust cols
for col in range(1, len(df.columns) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    ws.column_dimensions[get_column_letter(col)].width = max(15, len(str(cell.value)) + 2)

# filter
ws.auto_filter.ref = ws.dimensions

wb.save(file_path)
print("done")
